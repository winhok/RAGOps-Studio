#!/usr/bin/env python3
"""Create a synthetic Word upload for the document-to-citation walkthrough."""
import argparse
from pathlib import Path

from docx import Document


def create_document(output: Path):
    document = Document()
    document.add_heading('Document handling standard', level=1)
    document.add_paragraph('Synthetic data for RAGOps Studio demonstrations. These rules describe no real organization.')
    document.add_heading('Retention periods', level=2)
    document.add_paragraph('Preserve the original upload and its published revision.', style='List Bullet')
    document.add_paragraph('Apply the retention period listed below.', style='List Number')
    document.add_paragraph('Remove expired uploads after checking their current status.', style='List Number')
    table = document.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ('Record type', 'Retention period', 'Action'), strict=True):
        cell.text = text
    for values in [('Demonstration uploads', '30 days', 'Remove after expiry'), ('Demonstration audit reports', '90 days', 'Review before removal')]:
        for cell, text in zip(table.add_row().cells, values, strict=True):
            cell.text = text
    document.add_heading('Exceptions', level=2)
    document.add_paragraph('There is no policy in this demonstration document for retaining customer contracts.')
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.exists():
        raise FileExistsError('Choose a new output path; an existing file will not be overwritten')
    document.save(output)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('output', type=Path, help='New .docx file to create')
    args = parser.parse_args()
    if args.output.suffix.lower() != '.docx':
        parser.error('Output must have the .docx extension')
    create_document(args.output)
    print(f'Created synthetic upload: {args.output}')

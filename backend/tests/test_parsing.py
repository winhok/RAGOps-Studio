from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from app.core.errors import AppError
from app.services.parsing import extract_text
from conftest import auth


def word_document():
    document = Document()
    document.add_heading('Document handling', level=1)
    document.add_paragraph('Synthetic demonstration policy.')
    document.add_heading('Retention', level=2)
    document.add_paragraph('Keep the original upload.', style='List Bullet')
    document.add_paragraph('Check its retention period.', style='List Number')
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Record type'
    table.cell(0, 1).text = 'Retention period'
    table.cell(1, 0).text = 'Demonstration uploads'
    table.cell(1, 1).text = '30 days'
    document.add_paragraph('After the table: remove expired uploads.')
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_word_structure_and_body_order_are_preserved():
    text = extract_text('policy.DOCX', word_document())
    assert text.startswith('# Document handling')
    assert '## Retention' in text
    assert '- Keep the original upload.' in text
    assert '1. Check its retention period.' in text
    assert '| Demonstration uploads | 30 days |' in text
    assert text.index('Synthetic') < text.index('## Retention') < text.index('| Record type') < text.index('After the table')


def test_word_localized_headings_and_table_cells_are_escaped():
    document = Document()
    style = document.styles.add_style('标题 2', WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph('导入规则', style=style)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = '类型', '说明'
    table.cell(1, 0).text, table.cell(1, 1).text = 'A|B', '第一行\n第二行'
    output = BytesIO()
    document.save(output)
    text = extract_text('policy.docx', output.getvalue())
    assert '## 导入规则' in text
    assert 'A\\|B' in text
    assert '第一行<br>第二行' in text


def test_large_table_chunks_keep_headers_and_complete_rows(runtime, admin, publish):
    rows = [f'| Record-{i:03} | Retain for {i + 1} days |' for i in range(45)]
    publish(content='# Retention\n\n| Record | Period |\n| --- | --- |\n' + '\n'.join(rows), evidence_type='general')
    chunks = runtime.store.snapshot(admin)
    assert len(chunks) > 1
    assert all(len(c.text) <= 700 and '| Record | Period |\n| --- | --- |' in c.text for c in chunks)
    for row in rows:
        assert sum(row in c.text for c in chunks) == 1


def test_oversized_table_row_fails_without_replacing_active_version(runtime, admin, publish):
    publish()
    with pytest.raises(AppError, match='Table row'):
        publish(content='| Record | Notes |\n| --- | --- |\n| Upload | ' + 'x' * 900 + ' |', expected_version=1)
    assert runtime.store.source(admin, 'policy', 1).active
    assert len(runtime.store.versions(admin, 'policy')) == 1


def test_table_shaped_code_is_not_rejected_as_an_oversized_table(runtime, admin, publish):
    publish(content='# Code example\n\n```text\n| Record | Notes |\n| --- | --- |\n| Upload | ' + 'x' * 900 + ' |\n```')
    chunks = runtime.store.snapshot(admin)
    assert len(chunks) > 1 and all(len(c.text) <= 700 for c in chunks)


def test_word_upload_retrieves_table_value_with_bound_revision(client):
    response = client.post('/api/documents/upload', headers=auth(),
        data={'title': 'Document retention standard', 'evidence_type': 'general'},
        files={'file': ('policy.docx', word_document(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
    assert response.status_code == 200, response.text
    document = response.json()['document']
    answer = client.post('/api/chat', headers=auth(), json={'query': 'What is the document retention period?'}).json()
    assert answer['outcome'] == 'answered'
    assert '30 days' in answer['text']
    citation = next(c for c in answer['citations'] if c['document_id'] == document['id'] and '30 days' in c['content'])
    assert citation['version'] == 1
    source = client.get(citation['source_url'], headers=auth()).json()
    assert '| Demonstration uploads | 30 days |' in source['content']


@pytest.mark.parametrize('content', [b'not a Word file', b''])
def test_invalid_word_files_are_rejected(content):
    with pytest.raises(AppError):
        extract_text('policy.docx', content)


def test_word_expansion_limit_is_checked_before_parsing():
    output = BytesIO()
    with ZipFile(output, 'w', ZIP_DEFLATED) as archive:
        archive.writestr('word/document.xml', b'x' * 20000001)
    with pytest.raises(AppError, match='size limit'):
        extract_text('policy.docx', output.getvalue())
